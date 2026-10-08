"""Write :class:`PageDocument` to DOCX bytes."""

from __future__ import annotations

import io
from typing import BinaryIO

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt

from djangocms_tacc_export_page.document import Block, InlineSpan, PageDocument

# Monospace metadata (Word: Macro Text; fallback font for other editors)
_METADATA_STYLE = 'Macro Text'
_METADATA_FONT = 'Courier New'
_METADATA_SIZE_PT = 10
_CODE_FONT = 'Courier New'
_LINK_COLOR = '0563C1'


def write_page_document(document: PageDocument, stream: BinaryIO | None = None) -> bytes:
    """Serialize ``document`` to DOCX. Returns bytes when ``stream`` is omitted."""
    doc = Document()
    _add_title_block(doc, document)

    for block in document.blocks:
        _add_block(doc, block)

    if stream is not None:
        doc.save(stream)
        stream.seek(0)
        return stream.read()

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _add_metadata_paragraph(doc: Document, line: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    try:
        paragraph.style = doc.styles[_METADATA_STYLE]
    except KeyError:
        pass
    run = paragraph.add_run(line)
    run.font.name = _METADATA_FONT
    run.font.size = Pt(_METADATA_SIZE_PT)


def _add_title_block(doc: Document, document: PageDocument) -> None:
    title = document.title.strip() or 'Untitled page'
    meta_lines = [f'Title: {title}']
    if document.slug:
        meta_lines.append(f'Slug: {document.slug}')
    if document.page_url:
        meta_lines.append(f'URL: {document.page_url}')
    for line in meta_lines:
        _add_metadata_paragraph(doc, line)

    doc.add_paragraph('')


def _add_hyperlink(paragraph, text: str, url: str, span: InlineSpan) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')
    if span.bold:
        r_pr.append(OxmlElement('w:b'))
    if span.italic:
        r_pr.append(OxmlElement('w:i'))
    if span.code:
        r_fonts = OxmlElement('w:rFonts')
        r_fonts.set(qn('w:ascii'), _CODE_FONT)
        r_fonts.set(qn('w:hAnsi'), _CODE_FONT)
        r_pr.append(r_fonts)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    r_pr.append(underline)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), _LINK_COLOR)
    r_pr.append(color)
    run.append(r_pr)

    text_element = OxmlElement('w:t')
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_span_run(paragraph, span: InlineSpan) -> None:
    if span.url:
        _add_hyperlink(paragraph, span.text, span.url, span)
        return
    run = paragraph.add_run(span.text)
    if span.bold:
        run.bold = True
    if span.italic:
        run.italic = True
    if span.code:
        run.font.name = _CODE_FONT


def _add_rich_paragraph(doc: Document, block: Block, *, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    spans = block.runs or (InlineSpan(block.text),)
    for span in spans:
        if span.text:
            _add_span_run(paragraph, span)


def _add_block(doc: Document, block: Block) -> None:
    if block.kind == 'heading1':
        doc.add_heading(block.text, level=1)
        return
    if block.kind == 'heading2':
        doc.add_heading(block.text, level=2)
        return
    if block.kind == 'heading3':
        doc.add_heading(block.text, level=3)
        return
    if block.kind == 'bullet':
        _add_rich_paragraph(doc, block, style='List Bullet')
        return
    if block.kind == 'link_line':
        paragraph = doc.add_paragraph()
        if block.url:
            _add_hyperlink(
                paragraph,
                block.text,
                block.url,
                InlineSpan(block.text, url=block.url),
            )
        else:
            paragraph.add_run(block.text)
        return
    _add_rich_paragraph(doc, block)
