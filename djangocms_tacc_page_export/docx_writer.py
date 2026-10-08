"""Write :class:`PageDocument` to DOCX bytes."""

from __future__ import annotations

import io
from typing import BinaryIO

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

from djangocms_tacc_page_export.document import Block, PageDocument

# Monospace metadata (Word: Macro Text; fallback font for other editors)
_METADATA_STYLE = 'Macro Text'
_METADATA_FONT = 'Courier New'
_METADATA_SIZE_PT = 10


def write_page_document(document: PageDocument, stream: BinaryIO | None = None) -> bytes:
    """Serialize ``document`` to DOCX. Returns bytes when ``stream`` is omitted."""
    doc = Document()
    _add_title_block(doc, document)

    for block in document.blocks:
        _add_block(doc, block)

    if stream is not None:
        doc.save(stream)
        stream.seek(0)
        return stream.read()

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _add_metadata_paragraph(doc: Document, line: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    try:
        paragraph.style = doc.styles[_METADATA_STYLE]
    except KeyError:
        pass
    run = paragraph.add_run(line)
    run.font.name = _METADATA_FONT
    run.font.size = Pt(_METADATA_SIZE_PT)


def _add_title_block(doc: Document, document: PageDocument) -> None:
    title = document.title.strip() or 'Untitled page'
    meta_lines = [f'Title: {title}']
    if document.slug:
        meta_lines.append(f'Slug: {document.slug}')
    if document.page_url:
        meta_lines.append(f'URL: {document.page_url}')
    for line in meta_lines:
        _add_metadata_paragraph(doc, line)

    doc.add_paragraph('')


def _add_block(doc: Document, block: Block) -> None:
    if block.kind == 'heading1':
        doc.add_heading(block.text, level=1)
        return
    if block.kind == 'heading2':
        doc.add_heading(block.text, level=2)
        return
    if block.kind == 'heading3':
        doc.add_heading(block.text, level=3)
        return
    if block.kind == 'bullet':
        doc.add_paragraph(block.text, style='List Bullet')
        return
    if block.kind == 'link_line':
        paragraph = doc.add_paragraph()
        if block.url:
            paragraph.add_run(f'{block.text} ({block.url})')
        else:
            paragraph.add_run(block.text)
        return
    doc.add_paragraph(block.text)
